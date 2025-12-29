"""
Tamil Text Spacing Corrector for Quran Excel Files - V3
========================================================
Enhanced version with Surah/Verse location display for errors.

Features:
- Shows Surah and Verse numbers for each error found
- Displays before/after comparison
- Fixes Tamil spacing issues automatically
- Preserves original file, creates corrected copy

Author: Created for Spiral Nineteen IT Technologies
"""

import re
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import traceback

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from pathlib import Path


# Tamil Unicode ranges and characters
TAMIL_VOWEL_SIGNS = '\u0BBE\u0BBF\u0BC0\u0BC1\u0BC2\u0BC6\u0BC7\u0BC8\u0BCA\u0BCB\u0BCC'
TAMIL_PULLI = '\u0BCD'
TAMIL_ANUSVARA = '\u0B82'
TAMIL_DEPENDENT_MARKS = TAMIL_VOWEL_SIGNS + TAMIL_PULLI + TAMIL_ANUSVARA


def fix_tamil_spacing(text):
    """Fix spacing issues in Tamil text."""
    if not isinstance(text, str) or not text.strip():
        return text
    
    # Remove space(s) before Tamil dependent vowel signs and pulli
    pattern = r'\s+([' + TAMIL_DEPENDENT_MARKS + '])'
    text = re.sub(pattern, r'\1', text)
    
    # Remove zero-width characters
    text = re.sub(r'[\u200B\u200C\u200D\uFEFF]+', '', text)
    
    # Normalize multiple spaces to single space
    text = re.sub(r' {2,}', ' ', text)
    
    return text.strip()


def detect_tamil_issues(text):
    """Detect Tamil spacing issues in text."""
    issues = []
    if not isinstance(text, str):
        return issues
    
    # Check for space before dependent marks
    pattern = r'.{0,5}\s+[' + TAMIL_DEPENDENT_MARKS + '].{0,5}'
    matches = re.findall(pattern, text)
    for match in matches:
        issues.append(f"Space before matra: ...{match.strip()}...")
    
    if re.search(r' {2,}', text):
        issues.append("Multiple consecutive spaces")
    
    if re.search(r'[\u200B\u200C\u200D\uFEFF]', text):
        issues.append("Zero-width characters")
    
    return issues


def has_tamil_text(text):
    """Check if text contains Tamil characters."""
    if not isinstance(text, str):
        return False
    return bool(re.search(r'[\u0B80-\u0BFF]', text))


def find_surah_verse_columns(df):
    """
    Auto-detect Surah and Verse/Ayah columns in the dataframe.
    Returns tuple: (surah_col, verse_col) or (None, None) if not found.
    """
    surah_col = None
    verse_col = None
    
    # Common column name patterns for Surah
    surah_patterns = ['surah', 'sura', 'chapter', 'sūrah', 'sūra', 'surano', 'sura_no', 'surah_no', 'chapterno']
    
    # Common column name patterns for Verse/Ayah
    verse_patterns = ['verse', 'ayah', 'ayat', 'aya', 'āyah', 'verseno', 'verse_no', 'ayah_no', 'ayatno', 'ayano']
    
    for col in df.columns:
        col_lower = str(col).lower().replace(' ', '').replace('_', '')
        
        # Check for Surah column
        if surah_col is None:
            for pattern in surah_patterns:
                if pattern.replace('_', '') in col_lower:
                    surah_col = col
                    break
        
        # Check for Verse column
        if verse_col is None:
            for pattern in verse_patterns:
                if pattern.replace('_', '') in col_lower:
                    verse_col = col
                    break
    
    return surah_col, verse_col


class TamilFixerApp:
    """GUI Application for Tamil Text Fixer - Enhanced Version"""
    
    def __init__(self, root):
        self.root = root
        self.root.title("Tamil Quran Text Corrector V3 - Spiral Nineteen")
        self.root.geometry("850x650")
        self.root.resizable(True, True)
        
        self.setup_ui()
        self.input_file = None
        self.df = None
        self.surah_col = None
        self.verse_col = None
        self.tamil_columns = []
        
        self.check_dependencies()
    
    def check_dependencies(self):
        """Check if required libraries are installed."""
        missing = []
        if not PANDAS_AVAILABLE:
            missing.append("pandas")
        if not OPENPYXL_AVAILABLE:
            missing.append("openpyxl")
        
        if missing:
            msg = f"Missing libraries: {', '.join(missing)}\n\nPlease install using:\npip install {' '.join(missing)}"
            messagebox.showerror("Missing Dependencies", msg)
    
    def setup_ui(self):
        # Main frame
        main_frame = ttk.Frame(self.root, padding="15")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Title
        title_frame = ttk.Frame(main_frame)
        title_frame.pack(fill=tk.X, pady=(0, 10))
        
        title_label = ttk.Label(
            title_frame, 
            text="Tamil Quran Text Corrector",
            font=('Helvetica', 16, 'bold')
        )
        title_label.pack()
        
        subtitle_label = ttk.Label(
            title_frame,
            text="بِسْمِ اللَّهِ الرَّحْمَٰنِ الرَّحِيمِ",
            font=('Arial', 12)
        )
        subtitle_label.pack(pady=(5, 0))
        
        # File selection frame
        file_frame = ttk.LabelFrame(main_frame, text="Step 1: Select Excel File", padding="10")
        file_frame.pack(fill=tk.X, pady=(0, 10))
        
        self.file_label = ttk.Label(file_frame, text="No file selected", font=('Segoe UI', 9))
        self.file_label.pack(side=tk.LEFT, fill=tk.X, expand=True)
        
        browse_btn = ttk.Button(file_frame, text="Browse...", command=self.browse_file)
        browse_btn.pack(side=tk.RIGHT, padx=(10, 0))
        
        # Column Detection Frame
        col_frame = ttk.LabelFrame(main_frame, text="Step 2: Column Configuration", padding="10")
        col_frame.pack(fill=tk.X, pady=(0, 10))
        
        # Surah column
        row1 = ttk.Frame(col_frame)
        row1.pack(fill=tk.X, pady=(0, 5))
        ttk.Label(row1, text="Surah Column:", width=15).pack(side=tk.LEFT)
        self.surah_combo = ttk.Combobox(row1, width=25, state="readonly")
        self.surah_combo.pack(side=tk.LEFT, padx=(5, 10))
        ttk.Label(row1, text="(auto-detected)", foreground="gray").pack(side=tk.LEFT)
        
        # Verse column
        row2 = ttk.Frame(col_frame)
        row2.pack(fill=tk.X, pady=(0, 5))
        ttk.Label(row2, text="Verse/Ayah Column:", width=15).pack(side=tk.LEFT)
        self.verse_combo = ttk.Combobox(row2, width=25, state="readonly")
        self.verse_combo.pack(side=tk.LEFT, padx=(5, 10))
        ttk.Label(row2, text="(auto-detected)", foreground="gray").pack(side=tk.LEFT)
        
        # Tamil column
        row3 = ttk.Frame(col_frame)
        row3.pack(fill=tk.X)
        ttk.Label(row3, text="Tamil Column:", width=15).pack(side=tk.LEFT)
        self.tamil_combo = ttk.Combobox(row3, width=25, state="readonly")
        self.tamil_combo.pack(side=tk.LEFT, padx=(5, 10))
        ttk.Label(row3, text="(auto-detected)", foreground="gray").pack(side=tk.LEFT)
        
        # Progress frame
        progress_frame = ttk.Frame(main_frame)
        progress_frame.pack(fill=tk.X, pady=(0, 10))
        
        self.progress_label = ttk.Label(progress_frame, text="Ready", foreground="blue", font=('Segoe UI', 9, 'bold'))
        self.progress_label.pack(anchor=tk.W)
        
        self.progress_bar = ttk.Progressbar(progress_frame, mode='determinate', length=400)
        self.progress_bar.pack(fill=tk.X, pady=(5, 0))
        
        # Buttons frame
        btn_frame = ttk.Frame(main_frame)
        btn_frame.pack(fill=tk.X, pady=(5, 10))
        
        self.preview_btn = ttk.Button(
            btn_frame,
            text="🔍 Find Errors (Preview)",
            command=self.preview_issues,
            width=25
        )
        self.preview_btn.pack(side=tk.LEFT)
        
        self.process_btn = ttk.Button(
            btn_frame, 
            text="✓ Fix All Errors",
            command=self.process_file,
            width=20
        )
        self.process_btn.pack(side=tk.LEFT, padx=(10, 0))
        
        self.export_btn = ttk.Button(
            btn_frame,
            text="📋 Export Error Report",
            command=self.export_error_report,
            width=20
        )
        self.export_btn.pack(side=tk.LEFT, padx=(10, 0))
        
        clear_btn = ttk.Button(btn_frame, text="Clear Log", command=self.clear_log)
        clear_btn.pack(side=tk.RIGHT)
        
        # Results text area with better formatting
        results_frame = ttk.LabelFrame(main_frame, text="Step 3: Error Report (Surah:Verse Location)", padding="10")
        results_frame.pack(fill=tk.BOTH, expand=True)
        
        # Create text widget with scrollbar
        text_frame = ttk.Frame(results_frame)
        text_frame.pack(fill=tk.BOTH, expand=True)
        
        self.results_text = tk.Text(
            text_frame, 
            height=15, 
            wrap=tk.WORD, 
            font=('Consolas', 10),
            bg='#1e1e1e',
            fg='#d4d4d4',
            insertbackground='white'
        )
        scrollbar = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self.results_text.yview)
        
        self.results_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        self.results_text.config(yscrollcommand=scrollbar.set)
        
        # Configure text tags for colored output
        self.results_text.tag_configure('header', foreground='#569cd6', font=('Consolas', 10, 'bold'))
        self.results_text.tag_configure('success', foreground='#4ec9b0')
        self.results_text.tag_configure('error', foreground='#f14c4c')
        self.results_text.tag_configure('warning', foreground='#dcdcaa')
        self.results_text.tag_configure('location', foreground='#ce9178', font=('Consolas', 10, 'bold'))
        self.results_text.tag_configure('tamil', foreground='#c586c0')
        self.results_text.tag_configure('separator', foreground='#808080')
        
        # Initial message
        self.log("Application ready. Select an Excel file to begin.\n", 'success')
        
        # Store errors for export
        self.error_list = []
    
    def clear_log(self):
        self.results_text.delete(1.0, tk.END)
        self.error_list = []
    
    def log(self, message, tag=None):
        if tag:
            self.results_text.insert(tk.END, message, tag)
        else:
            self.results_text.insert(tk.END, message)
        self.results_text.see(tk.END)
        self.root.update_idletasks()
    
    def browse_file(self):
        file_path = filedialog.askopenfilename(
            title="Select Quran Excel File",
            filetypes=[
                ("Excel Files", "*.xlsx *.xls"),
                ("All Files", "*.*")
            ]
        )
        if file_path:
            self.input_file = file_path
            filename = Path(file_path).name
            self.file_label.config(text=filename)
            self.log(f"\n✓ Selected: {filename}\n", 'success')
            
            # Load file and detect columns
            self.load_and_detect_columns()
    
    def load_and_detect_columns(self):
        """Load Excel file and auto-detect relevant columns."""
        try:
            self.update_progress("Loading file...", 10)
            self.df = pd.read_excel(self.input_file)
            
            columns = list(self.df.columns)
            self.log(f"  Columns found: {columns}\n")
            
            # Update comboboxes with column options
            self.surah_combo['values'] = ['(None)'] + columns
            self.verse_combo['values'] = ['(None)'] + columns
            self.tamil_combo['values'] = ['(None)'] + columns
            
            # Auto-detect Surah and Verse columns
            self.surah_col, self.verse_col = find_surah_verse_columns(self.df)
            
            if self.surah_col:
                self.surah_combo.set(self.surah_col)
                self.log(f"  ✓ Surah column detected: '{self.surah_col}'\n", 'success')
            else:
                self.surah_combo.set('(None)')
                self.log(f"  ⚠ Surah column not auto-detected. Please select manually.\n", 'warning')
            
            if self.verse_col:
                self.verse_combo.set(self.verse_col)
                self.log(f"  ✓ Verse column detected: '{self.verse_col}'\n", 'success')
            else:
                self.verse_combo.set('(None)')
                self.log(f"  ⚠ Verse column not auto-detected. Please select manually.\n", 'warning')
            
            # Auto-detect Tamil columns
            self.tamil_columns = []
            for col in columns:
                if 'tamil' in str(col).lower():
                    self.tamil_columns.append(col)
                else:
                    sample = self.df[col].dropna().head(20)
                    if any(has_tamil_text(str(val)) for val in sample):
                        self.tamil_columns.append(col)
            
            if self.tamil_columns:
                self.tamil_combo.set(self.tamil_columns[0])
                self.log(f"  ✓ Tamil column detected: '{self.tamil_columns[0]}'\n", 'success')
            else:
                self.tamil_combo.set('(None)')
                self.log(f"  ⚠ Tamil column not detected. Please select manually.\n", 'warning')
            
            self.update_progress("Ready", 0)
            self.log(f"\n→ Click 'Find Errors' to scan for Tamil spacing issues.\n", 'header')
            
        except Exception as e:
            self.log(f"\n❌ Error loading file: {str(e)}\n", 'error')
            messagebox.showerror("Error", f"Failed to load file:\n{str(e)}")
    
    def update_progress(self, message, value):
        self.progress_label.config(text=message)
        self.progress_bar['value'] = value
        self.root.update_idletasks()
    
    def get_location_string(self, idx):
        """Get Surah:Verse location string for a given row index."""
        surah_col = self.surah_combo.get()
        verse_col = self.verse_combo.get()
        
        surah = ""
        verse = ""
        
        if surah_col and surah_col != '(None)':
            surah_val = self.df.at[idx, surah_col]
            surah = str(int(surah_val)) if pd.notna(surah_val) else "?"
        
        if verse_col and verse_col != '(None)':
            verse_val = self.df.at[idx, verse_col]
            verse = str(int(verse_val)) if pd.notna(verse_val) else "?"
        
        if surah and verse:
            return f"Surah {surah}, Verse {verse}"
        elif surah:
            return f"Surah {surah}"
        elif verse:
            return f"Verse {verse}"
        else:
            return f"Row {idx + 2}"  # Excel row number (1-indexed + header)
    
    def preview_issues(self):
        """Preview Tamil spacing issues with Surah/Verse locations."""
        if self.df is None:
            messagebox.showwarning("No File", "Please select an Excel file first.")
            return
        
        tamil_col = self.tamil_combo.get()
        if not tamil_col or tamil_col == '(None)':
            messagebox.showwarning("No Tamil Column", "Please select the Tamil column.")
            return
        
        self.clear_log()
        self.error_list = []
        
        self.log("═" * 60 + "\n", 'separator')
        self.log("  TAMIL SPACING ERROR REPORT\n", 'header')
        self.log("═" * 60 + "\n\n", 'separator')
        
        self.update_progress("Scanning for errors...", 10)
        self.preview_btn.config(state=tk.DISABLED)
        
        try:
            total_rows = len(self.df)
            error_count = 0
            
            self.log(f"Scanning column: '{tamil_col}'\n", 'header')
            self.log(f"Total rows: {total_rows}\n\n")
            self.log("-" * 60 + "\n", 'separator')
            
            for idx in range(total_rows):
                val = self.df.at[idx, tamil_col]
                
                if isinstance(val, str):
                    issues = detect_tamil_issues(val)
                    
                    if issues:
                        error_count += 1
                        location = self.get_location_string(idx)
                        
                        # Get before/after
                        original = val
                        fixed = fix_tamil_spacing(val)
                        
                        # Store for export
                        surah_col = self.surah_combo.get()
                        verse_col = self.verse_combo.get()
                        surah_val = self.df.at[idx, surah_col] if surah_col != '(None)' else ''
                        verse_val = self.df.at[idx, verse_col] if verse_col != '(None)' else ''
                        
                        self.error_list.append({
                            'row': idx + 2,
                            'surah': surah_val,
                            'verse': verse_val,
                            'location': location,
                            'issue': issues[0],
                            'before': original,
                            'after': fixed
                        })
                        
                        # Display error
                        self.log(f"\n📍 ", 'warning')
                        self.log(f"{location}\n", 'location')
                        self.log(f"   Issue: {issues[0]}\n", 'error')
                        
                        # Show before/after (truncated if too long)
                        before_display = original[:80] + "..." if len(original) > 80 else original
                        after_display = fixed[:80] + "..." if len(fixed) > 80 else fixed
                        
                        self.log(f"   Before: ", 'warning')
                        self.log(f"{before_display}\n", 'tamil')
                        self.log(f"   After:  ", 'success')
                        self.log(f"{after_display}\n", 'tamil')
                
                # Update progress
                if idx % 200 == 0:
                    progress = 10 + int((idx / total_rows) * 80)
                    self.update_progress(f"Scanning... {idx}/{total_rows}", progress)
            
            # Summary
            self.log("\n" + "═" * 60 + "\n", 'separator')
            self.log("  SUMMARY\n", 'header')
            self.log("═" * 60 + "\n", 'separator')
            self.log(f"\n  Total rows scanned: {total_rows}\n")
            self.log(f"  Errors found: ", 'warning')
            self.log(f"{error_count}\n", 'error' if error_count > 0 else 'success')
            
            if error_count > 0:
                self.log(f"\n  → Click 'Fix All Errors' to correct these issues.\n", 'header')
                self.log(f"  → Click 'Export Error Report' to save the error list.\n", 'header')
            else:
                self.log(f"\n  ✓ No Tamil spacing errors found! Alhamdulillah!\n", 'success')
            
            self.update_progress(f"Found {error_count} errors", 100)
            
        except Exception as e:
            self.log(f"\n❌ Error: {str(e)}\n", 'error')
            self.log(traceback.format_exc(), 'error')
            messagebox.showerror("Error", str(e))
        
        finally:
            self.preview_btn.config(state=tk.NORMAL)
    
    def export_error_report(self):
        """Export error list to Excel file."""
        if not self.error_list:
            messagebox.showinfo("No Errors", "No errors to export. Run 'Find Errors' first.")
            return
        
        try:
            # Ask for save location
            output_path = filedialog.asksaveasfilename(
                title="Save Error Report",
                defaultextension=".xlsx",
                filetypes=[("Excel Files", "*.xlsx")],
                initialfilename="tamil_error_report.xlsx"
            )
            
            if not output_path:
                return
            
            # Create DataFrame from error list
            report_df = pd.DataFrame(self.error_list)
            report_df.columns = ['Excel Row', 'Surah', 'Verse', 'Location', 'Issue Type', 'Before (Original)', 'After (Fixed)']
            
            # Save to Excel
            report_df.to_excel(output_path, index=False)
            
            self.log(f"\n✓ Error report exported to: {output_path}\n", 'success')
            messagebox.showinfo("Exported", f"Error report saved to:\n{Path(output_path).name}")
            
        except Exception as e:
            messagebox.showerror("Export Error", str(e))
    
    def process_file(self):
        """Process and fix Tamil text in the file."""
        if self.df is None:
            messagebox.showwarning("No File", "Please select an Excel file first.")
            return
        
        tamil_col = self.tamil_combo.get()
        if not tamil_col or tamil_col == '(None)':
            messagebox.showwarning("No Tamil Column", "Please select the Tamil column.")
            return
        
        # Confirm
        if not messagebox.askyesno("Confirm", f"This will fix all Tamil spacing errors in column '{tamil_col}'.\n\nContinue?"):
            return
        
        self.log("\n" + "═" * 60 + "\n", 'separator')
        self.log("  PROCESSING FILE...\n", 'header')
        self.log("═" * 60 + "\n", 'separator')
        
        self.update_progress("Fixing errors...", 10)
        self.process_btn.config(state=tk.DISABLED)
        self.preview_btn.config(state=tk.DISABLED)
        
        try:
            total_rows = len(self.df)
            fixes = 0
            
            for idx in range(total_rows):
                original = self.df.at[idx, tamil_col]
                if isinstance(original, str):
                    fixed = fix_tamil_spacing(original)
                    if fixed != original:
                        self.df.at[idx, tamil_col] = fixed
                        fixes += 1
                
                if idx % 200 == 0:
                    progress = 10 + int((idx / total_rows) * 70)
                    self.update_progress(f"Fixing... {idx}/{total_rows}", progress)
            
            # Save file
            self.update_progress("Saving file...", 85)
            
            input_path = Path(self.input_file)
            output_path = input_path.parent / f"{input_path.stem}_corrected{input_path.suffix}"
            
            self.df.to_excel(output_path, index=False)
            
            self.update_progress("Complete!", 100)
            
            self.log(f"\n✓ Fixed {fixes} rows\n", 'success')
            self.log(f"✓ Saved to: {output_path.name}\n", 'success')
            self.log(f"\nAlhamdulillah! Processing complete.\n", 'header')
            
            messagebox.showinfo(
                "Success - Alhamdulillah!",
                f"Fixed {fixes} rows!\n\nSaved to:\n{output_path.name}"
            )
            
        except Exception as e:
            self.log(f"\n❌ Error: {str(e)}\n", 'error')
            messagebox.showerror("Error", str(e))
        
        finally:
            self.process_btn.config(state=tk.NORMAL)
            self.preview_btn.config(state=tk.NORMAL)


def main():
    root = tk.Tk()
    app = TamilFixerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
